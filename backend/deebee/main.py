from __future__ import annotations

import asyncio
import csv
import io
import json
import secrets
from typing import Any, Literal

from fastapi import Depends, FastAPI, File, Header, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from .config import settings
from .jobs import jobs
from .mysql import DeeBeeError
from .security import issue_token, verify_token
from .workbench import workbench


app = FastAPI(title="DeeBee API", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=list(settings.cors_origins),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(DeeBeeError)
async def deebee_error_handler(_, exc: DeeBeeError) -> JSONResponse:
    return JSONResponse(
        status_code=400,
        content={"error": {"message": str(exc), "code": exc.code}},
    )


class LoginBody(BaseModel):
    username: str
    password: str


class ConnectionBaseBody(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)

    driver: Literal["mysql", "postgresql"]
    name: str = Field(min_length=1, max_length=100)
    host: str = Field(min_length=1, max_length=255)
    port: int = Field(ge=1, le=65535)
    user: str = Field(min_length=1, max_length=255)
    default_database: str = Field(default="", max_length=255)
    default_schema: str = Field(default="", max_length=255)


class ConnectionCreateBody(ConnectionBaseBody):
    password: str = Field(default="", max_length=4096)


class ConnectionUpdateBody(ConnectionBaseBody):
    password: str | None = Field(default=None, max_length=4096)


class ConnectionTestBody(ConnectionUpdateBody):
    profile_id: str = Field(default="", max_length=100)


class SchemaBody(BaseModel):
    schema_: str = ""

    @model_validator(mode="before")
    @classmethod
    def accept_schema_name(cls, value: Any) -> Any:
        if isinstance(value, dict) and "schema" in value and "schema_" not in value:
            value = {**value, "schema_": value["schema"]}
        return value


class SessionBody(SchemaBody):
    profile_id: str = "mysql-default"
    database: str = ""
    autocommit: bool = True
    workspace_id: str = Field(default="", max_length=100)


class WorkspaceBody(BaseModel):
    workspace_id: str = Field(min_length=1, max_length=100)


class AutocommitBody(BaseModel):
    enabled: bool


class QueryBody(BaseModel):
    sql: str
    limit: int = Field(default=1000, ge=1, le=10000)


class FilterItem(BaseModel):
    column: str
    operator: Literal["eq", "ne", "gt", "gte", "lt", "lte", "contains", "is_null", "not_null"] = "eq"
    value: Any = None


class SortItem(BaseModel):
    column: str
    direction: Literal["asc", "desc"] = "asc"


class DataRequest(SchemaBody):
    profile_id: str = "mysql-default"
    database: str
    table: str
    page: int = Field(default=1, ge=1)
    page_size: int = Field(default=100, ge=1, le=1000)
    filters: list[FilterItem] = Field(default_factory=list)
    sort: SortItem | None = None


class RowInsert(SchemaBody):
    profile_id: str = "mysql-default"
    database: str
    table: str
    values: dict[str, Any]


class RowUpdate(SchemaBody):
    profile_id: str = "mysql-default"
    database: str
    table: str
    key: dict[str, Any]
    changes: dict[str, Any]


class RowDelete(SchemaBody):
    profile_id: str = "mysql-default"
    database: str
    table: str
    key: dict[str, Any]


class DdlPreviewBody(BaseModel):
    profile_id: str = "mysql-default"
    spec: dict[str, Any]
    current_table: str | None = None


class DdlApplyBody(DdlPreviewBody):
    expected_statements: list[str]


class DatabaseCreateBody(BaseModel):
    profile_id: str = "mysql-default"
    name: str
    charset: str = "utf8mb4"
    collation: str = "utf8mb4_unicode_ci"


class ObjectActionBody(SchemaBody):
    profile_id: str = "mysql-default"
    database: str
    table: str
    action: Literal["drop", "empty", "truncate", "rename", "duplicate", "check", "analyze", "optimize", "repair"]
    target: str = ""
    with_data: bool = False


class GenerateBody(SchemaBody):
    profile_id: str = "mysql-default"
    database: str
    table: str
    count: int = Field(default=100, ge=1, le=10000)


class GenericObjectActionBody(SchemaBody):
    profile_id: str = "mysql-default"
    database: str
    kind: Literal["view", "materialized view", "function", "procedure", "trigger", "event"]
    name: str
    action: Literal["drop", "enable", "disable"]


def current_user(authorization: str | None = Header(default=None)) -> str:
    if not authorization or not authorization.startswith("Bearer "):
        raise DeeBeeError("登录已失效")
    user = verify_token(authorization.removeprefix("Bearer ").strip())
    if not user:
        raise DeeBeeError("登录已失效")
    return user


def _parse_import_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "csv":
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    if suffix == "json":
        parsed = json.loads(content)
        if not isinstance(parsed, list) or not all(isinstance(item, dict) for item in parsed):
            raise DeeBeeError("JSON 文件必须是对象数组")
        return parsed
    if suffix in {"xlsx", "xlsm"}:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            raise DeeBeeError("Excel 文件没有数据")
        headers = [str(item) if item is not None else "" for item in values[0]]
        return [dict(zip(headers, row)) for row in values[1:]]
    raise DeeBeeError("仅支持 CSV、JSON 和 XLSX 文件")


@app.get("/api/health")
async def health() -> dict[str, Any]:
    drivers = list(dict.fromkeys(item["driver"] for item in workbench.list_profiles()))
    return {
        "ok": True, "service": "deebee",
        "driver": drivers[0] if len(drivers) == 1 else "multi",
        "drivers": drivers,
    }


@app.post("/api/auth/login")
async def login(body: LoginBody) -> dict[str, Any]:
    valid_user = secrets.compare_digest(body.username, settings.admin_user)
    valid_password = secrets.compare_digest(body.password, settings.admin_password)
    if not (valid_user and valid_password):
        raise DeeBeeError("用户名或密码错误")
    return {"token": issue_token(body.username), "user": {"username": body.username}}


@app.get("/api/auth/me")
async def me(user: str = Depends(current_user)) -> dict[str, Any]:
    return {"username": user}


@app.get("/api/connections")
async def connections(_: str = Depends(current_user)) -> list[dict[str, Any]]:
    return workbench.list_profiles()


@app.post("/api/connections/test")
async def test_connection_settings(
    body: ConnectionTestBody, _: str = Depends(current_user)
) -> dict[str, Any]:
    values = body.model_dump(exclude={"profile_id"})
    return await asyncio.to_thread(workbench.test_connection, values, body.profile_id)


@app.post("/api/connections", status_code=201)
async def create_connection(
    body: ConnectionCreateBody, _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.create_connection, body.model_dump())


@app.patch("/api/connections/{profile_id}")
async def update_connection(
    profile_id: str, body: ConnectionUpdateBody, _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.update_connection, profile_id, body.model_dump()
    )


@app.delete("/api/connections/{profile_id}", status_code=204)
async def delete_connection(profile_id: str, _: str = Depends(current_user)) -> Response:
    await asyncio.to_thread(workbench.delete_connection, profile_id)
    return Response(status_code=204)


@app.post("/api/connections/{profile_id}/test")
async def test_connection(profile_id: str, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.test_profile, profile_id)


@app.get("/api/connections/{profile_id}/databases")
async def databases(profile_id: str, _: str = Depends(current_user)) -> list[dict[str, Any]]:
    return await asyncio.to_thread(workbench.databases, profile_id)


@app.get("/api/connections/{profile_id}/databases/{database}/schemas")
async def schemas(
    profile_id: str, database: str, _: str = Depends(current_user)
) -> list[dict[str, Any]]:
    return await asyncio.to_thread(workbench.schemas, profile_id, database)


@app.get("/api/connections/{profile_id}/databases/{database}/objects")
async def objects(
    profile_id: str, database: str, schema: str = "", _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.objects, profile_id, database, schema)


@app.get("/api/connections/{profile_id}/databases/{database}/catalog")
async def catalog(
    profile_id: str, database: str, schema: str = "", _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.catalog, profile_id, database, schema)


@app.post("/api/databases")
async def create_database(body: DatabaseCreateBody, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.create_database, body.profile_id, body.name, body.charset, body.collation
    )


@app.delete("/api/databases/{profile_id}/{database}")
async def drop_database(profile_id: str, database: str, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.drop_database, profile_id, database)


@app.post("/api/objects/table/action")
async def table_action(body: ObjectActionBody, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.table_action, body.profile_id, body.database, body.table, body.action,
        target=body.target, with_data=body.with_data, schema=body.schema_,
    )


@app.get("/api/objects/{profile_id}/{database}/{kind}/{name}/ddl")
async def object_ddl(
    profile_id: str, database: str, kind: str, name: str, schema: str = "",
    _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.object_ddl, profile_id, database, kind, name, schema)


@app.post("/api/objects/action")
async def generic_object_action(body: GenericObjectActionBody, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.object_action, body.profile_id, body.database, body.kind, body.name,
        body.action, body.schema_
    )


@app.get("/api/search/{profile_id}/{database}")
async def search_objects(
    profile_id: str, database: str, q: str = Query(min_length=1, max_length=200), schema: str = "",
    _: str = Depends(current_user),
) -> list[dict[str, Any]]:
    return await asyncio.to_thread(workbench.search_objects, profile_id, database, q, schema)


@app.get("/api/privileges/{profile_id}/{database}")
async def privileges(
    profile_id: str, database: str, table: str = "", schema: str = "", _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.grants, profile_id, database, table, schema)


@app.post("/api/data/generate")
async def generate_data(body: GenerateBody, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.generate_data, body.profile_id, body.database, body.table, body.count,
        body.schema_
    )


@app.post("/api/sessions")
async def create_session(body: SessionBody, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.create_session, body.profile_id, body.database, body.autocommit,
        body.workspace_id, body.schema_
    )


@app.post("/api/sessions/cleanup")
async def cleanup_sessions(body: WorkspaceBody, _: str = Depends(current_user)) -> dict[str, int]:
    closed = await asyncio.to_thread(workbench.cleanup_workspace, body.workspace_id)
    return {"closed": closed}


@app.get("/api/sessions/{session_id}")
async def inspect_session(session_id: str, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.inspect_session, session_id)


@app.delete("/api/sessions/{session_id}")
async def close_session(session_id: str, _: str = Depends(current_user)) -> dict[str, bool]:
    await asyncio.to_thread(workbench.close_session, session_id)
    return {"ok": True}


@app.patch("/api/sessions/{session_id}/autocommit")
async def set_autocommit(
    session_id: str, body: AutocommitBody, _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.set_autocommit, session_id, body.enabled)


@app.post("/api/sessions/{session_id}/query")
async def run_query(
    session_id: str, body: QueryBody, _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.execute, session_id, body.sql, body.limit)


@app.post("/api/sessions/{session_id}/cancel")
async def cancel_query(session_id: str, _: str = Depends(current_user)) -> dict[str, bool]:
    return {"cancelled": await asyncio.to_thread(workbench.cancel, session_id)}


@app.post("/api/sessions/{session_id}/commit")
async def commit(session_id: str, _: str = Depends(current_user)) -> dict[str, bool]:
    await asyncio.to_thread(workbench.commit, session_id)
    return {"ok": True}


@app.post("/api/sessions/{session_id}/rollback")
async def rollback(session_id: str, _: str = Depends(current_user)) -> dict[str, bool]:
    await asyncio.to_thread(workbench.rollback, session_id)
    return {"ok": True}


@app.get("/api/schema/{profile_id}/{database}/{table}")
async def table_schema(
    profile_id: str, database: str, table: str, schema: str = "", _: str = Depends(current_user)
) -> dict[str, Any]:
    return await asyncio.to_thread(workbench.table_schema, profile_id, database, table, schema)


@app.post("/api/data/read")
async def read_data(body: DataRequest, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.table_data,
        body.profile_id,
        body.database,
        body.table,
        body.page,
        body.page_size,
        [item.model_dump() for item in body.filters],
        body.sort.model_dump() if body.sort else None,
        body.schema_,
    )


@app.post("/api/data/rows")
async def insert_row(body: RowInsert, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.insert_row, body.profile_id, body.database, body.table, body.values,
        body.schema_
    )


@app.patch("/api/data/rows")
async def update_row(body: RowUpdate, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.update_row,
        body.profile_id,
        body.database,
        body.table,
        body.key,
        body.changes,
        body.schema_,
    )


@app.delete("/api/data/rows")
async def delete_row(body: RowDelete, _: str = Depends(current_user)) -> dict[str, Any]:
    return await asyncio.to_thread(
        workbench.delete_row, body.profile_id, body.database, body.table, body.key,
        body.schema_
    )


@app.get("/api/data/export")
async def export_data(
    profile_id: str,
    database: str,
    table: str,
    schema: str = "",
    format: Literal["csv", "json", "sql", "xlsx"] = Query(default="csv"),
    _: str = Depends(current_user),
) -> Response:
    payload = await asyncio.to_thread(
        workbench.table_data, profile_id, database, table, 1, 100000, [], None, schema
    )
    rows = payload["rows"]
    columns = [item["name"] for item in payload["columns"]]
    filename = f"{table}.{format}"
    if format == "json":
        data = json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8")
        media_type = "application/json"
    elif format == "sql":
        quote = '"' if workbench.profile_driver(profile_id) == "postgresql" else "`"
        def quote_name(value: str) -> str:
            return quote + value.replace(quote, quote + quote) + quote
        output = []
        for row in rows:
            values = []
            for column in columns:
                value = row.get(column)
                if value is None:
                    values.append("NULL")
                elif isinstance(value, (int, float)):
                    values.append(str(value))
                else:
                    escaped = str(value).replace("'", "''")
                    values.append("'" + escaped + "'")
            escaped_columns = ", ".join(quote_name(column) for column in columns)
            target = f"{quote_name(schema)}.{quote_name(table)}" if schema else quote_name(table)
            output.append(f"INSERT INTO {target} ({escaped_columns}) VALUES ({', '.join(values)});")
        data = ("\n".join(output) + "\n").encode("utf-8")
        media_type = "application/sql"
    elif format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = table[:31]
        sheet.append(columns)
        for row in rows:
            sheet.append([json.dumps(row.get(column), ensure_ascii=False) if isinstance(row.get(column), dict) else row.get(column) for column in columns])
        buffer = io.BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        text = io.StringIO()
        writer = csv.DictWriter(text, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
        data = ("\ufeff" + text.getvalue()).encode("utf-8")
        media_type = "text/csv"
    return Response(
        data,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/data/import")
async def import_data(
    profile_id: str,
    database: str,
    table: str,
    schema: str = "",
    file: UploadFile = File(...),
    _: str = Depends(current_user),
) -> dict[str, Any]:
    content = await file.read()
    rows = _parse_import_rows(file.filename or "", content)
    if len(rows) > 100000:
        raise DeeBeeError("单次最多导入 100000 行")
    return await asyncio.to_thread(workbench.bulk_insert, profile_id, database, table, rows, schema)


@app.post("/api/sql/execute-file")
async def execute_sql_file(
    profile_id: str,
    database: str,
    schema: str = "",
    file: UploadFile = File(...),
    _: str = Depends(current_user),
) -> dict[str, Any]:
    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise DeeBeeError("SQL 文件不能超过 50 MB")
    try:
        sql = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DeeBeeError("SQL 文件必须使用 UTF-8 编码") from exc
    return await asyncio.to_thread(workbench.execute_script, profile_id, database, sql, schema)


@app.post("/api/jobs/sql-file")
async def execute_sql_file_job(
    profile_id: str, database: str, schema: str = "", file: UploadFile = File(...),
    _: str = Depends(current_user),
) -> dict[str, Any]:
    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise DeeBeeError("SQL 文件不能超过 50 MB")
    try:
        sql = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DeeBeeError("SQL 文件必须使用 UTF-8 编码") from exc
    def task(progress, cancelled):
        progress(10, "正在解析 SQL")
        if cancelled.is_set(): return None
        progress(25, "正在执行 SQL")
        return workbench.execute_script(profile_id, database, sql, schema)
    return jobs.create("sql-file", task)


@app.post("/api/jobs/import")
async def import_data_job(
    profile_id: str, database: str, table: str, schema: str = "", file: UploadFile = File(...),
    _: str = Depends(current_user),
) -> dict[str, Any]:
    content = await file.read()
    filename = file.filename or ""
    def task(progress, cancelled):
        progress(10, "正在解析文件")
        rows = _parse_import_rows(filename, content)
        if len(rows) > 100000:
            raise DeeBeeError("单次最多导入 100000 行")
        if cancelled.is_set(): return None
        progress(45, f"正在写入 {len(rows)} 行")
        result = workbench.bulk_insert(profile_id, database, table, rows, schema)
        progress(95, "正在刷新表数据")
        return result
    return jobs.create("import", task)


@app.post("/api/jobs/generate")
async def generate_data_job(body: GenerateBody, _: str = Depends(current_user)) -> dict[str, Any]:
    def task(progress, cancelled):
        total = 0
        remaining = body.count
        while remaining and not cancelled.is_set():
            batch = min(500, remaining)
            result = workbench.generate_data(
                body.profile_id, body.database, body.table, batch, body.schema_
            )
            total += int(result.get("affected_rows", 0)); remaining -= batch
            progress(10 + int(85 * total / body.count), f"已生成 {total}/{body.count} 行")
        return {"affected_rows": total}
    return jobs.create("generate", task)


@app.get("/api/jobs/{job_id}")
async def job_status(job_id: str, _: str = Depends(current_user)) -> dict[str, Any]:
    try: return jobs.public(job_id)
    except KeyError as exc: raise DeeBeeError("任务不存在或已过期") from exc


@app.post("/api/jobs/{job_id}/cancel")
async def cancel_job(job_id: str, _: str = Depends(current_user)) -> dict[str, Any]:
    try: return jobs.cancel(job_id)
    except KeyError as exc: raise DeeBeeError("任务不存在或已过期") from exc


@app.get("/api/sql/dump")
async def dump_sql(
    profile_id: str,
    database: str,
    table: str = "",
    include_data: bool = True,
    schema: str = "",
    _: str = Depends(current_user),
) -> Response:
    sql = await asyncio.to_thread(
        workbench.dump_sql, profile_id, database, table, include_data, schema
    )
    filename = f"{table or database}.sql"
    return Response(
        sql.encode("utf-8"),
        media_type="application/sql",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/ddl/preview")
async def preview_ddl(body: DdlPreviewBody, _: str = Depends(current_user)) -> dict[str, Any]:
    statements = await asyncio.to_thread(
        workbench.preview_ddl, body.profile_id, body.spec, body.current_table
    )
    dangerous = [
        statement
        for statement in statements
        if " DROP " in statement.upper() or " MODIFY COLUMN " in statement.upper()
    ]
    return {"statements": statements, "dangerous": dangerous}


@app.post("/api/ddl/apply")
async def apply_ddl(body: DdlApplyBody, _: str = Depends(current_user)) -> dict[str, Any]:
    generated = await asyncio.to_thread(
        workbench.preview_ddl, body.profile_id, body.spec, body.current_table
    )
    if generated != body.expected_statements:
        raise DeeBeeError("表结构在预览后发生变化，请重新生成 DDL")
    database = str(body.spec.get("database", ""))
    return await asyncio.to_thread(workbench.apply_ddl, body.profile_id, database, generated)
